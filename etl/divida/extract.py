"""
Dívida pública federal — o que ela custa por ano, quanto é o estoque e de que ela é feita.

    python -m etl.divida.extract            # gera data/public/divida.json

Três números diferentes costumam ser chamados de "o gasto com a dívida", e a página separa os três:

1. **Juros nominais do Governo Central** — o custo econômico do ano, apurado pelo Tesouro. Vem da planilha
   *Resultado do Tesouro Nacional — Série Histórica*, aba **2.1**, linha **"9. JUROS NOMINAIS"**, em R$ milhões
   correntes, **desde 1997**. É o número mais próximo de "quanto a dívida custou neste ano".
2. **Juros e encargos pagos pelo Orçamento** — a despesa orçamentária da função 28, que a página lê de
   `poderes_categoria_<ano>.json` (SIOP). É menor que o item 1 porque boa parte dos juros é apropriada no estoque dos
   títulos e só vira pagamento no resgate.
3. **Amortização** — também está no Orçamento e é o maior valor de todos, mas **não é custo**: é rolagem, dívida
   velha paga com dívida nova. Somá-la aos juros é o erro mais comum nessa conta.

O **estoque** vem do arquivo *Estoque da Dívida Pública Federal* do Tesouro Transparente (CSV, um registro por título
e mês, desde set/2017), de onde saem o total, a divisão entre dívida interna e externa, o perfil de vencimento e a
composição por indexador — a parte que explica por que a conta reage rápido à Selic.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import httpx
import openpyxl

from etl.estatais.dividendos import _planilha as _planilha_rtn

log = logging.getLogger("etl.divida")
PUB, RAW = Path("data/public"), Path("data/raw")
CKAN_ESTOQUE = ("https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show"
                "?id=estoque-da-divida-publica-federal")
ABA_RTN, LINHA_JUROS = "2.1", re.compile(r"^\s*9\.\s*JUROS NOMINAIS", re.I)

# A que cada família de título é indexada. Definições do próprio Tesouro (Relatório Anual da Dívida Pública).
INDEXADOR = [
    ("LFT", "Selic (pós-fixado)"),
    ("NTN-B", "IPCA"),
    ("NTN-C", "IGP-M"),
    ("LTN", "Prefixado"),
    ("NTN-F", "Prefixado"),
]


def _num(v: str) -> float | None:
    try:
        return float((v or "").strip().replace(".", "").replace(",", "."))
    except ValueError:
        return None


def _ordem_mes(m: str) -> tuple[str, str]:
    return (m[3:], m[:2])                      # "08/2026" → ("2026", "08")


def juros_nominais(c: httpx.Client) -> dict[str, float]:
    """Série anual do custo da dívida para o Governo Central, desde 1997."""
    conteudo = _planilha_rtn(c)
    if conteudo is None:
        return {}
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    if ABA_RTN not in wb.sheetnames:
        log.warning("aba %s não existe na planilha do RTN", ABA_RTN)
        return {}
    linhas = list(wb[ABA_RTN].iter_rows(values_only=True))
    anos: list[str] = []
    for row in linhas[:12]:
        if str(row[0] or "").strip().lower().startswith("discrimina"):
            anos = [str(int(v)) for v in row[1:] if isinstance(v, (int, float)) and 1990 < v < 2100]
            break
    for row in linhas:
        if LINHA_JUROS.match(str(row[0] or "")):
            # a planilha publica despesa com sinal negativo; aqui vira valor positivo, em reais
            return {a: round(abs(row[i + 1]) * 1e6, 2) for i, a in enumerate(anos)
                    if i + 1 < len(row) and isinstance(row[i + 1], (int, float))}
    log.warning("linha de juros nominais não encontrada na aba %s", ABA_RTN)
    return {}


def estoque(c: httpx.Client) -> dict:
    """Estoque da DPF por mês, e a composição do mês mais recente."""
    destino = RAW / "estoque_dpf.csv"
    if not (destino.exists() and destino.stat().st_size > 1_000_000):
        r = c.get(CKAN_ESTOQUE)
        if r.status_code != 200:
            log.warning("CKAN do estoque respondeu HTTP %s", r.status_code)
            return {}
        url = next((x["url"] for x in (r.json().get("result") or {}).get("resources", [])
                    if (x.get("format") or "").upper() == "CSV"), None)
        if not url:
            log.warning("CSV do estoque não encontrado no CKAN")
            return {}
        d = c.get(url)
        if d.status_code != 200:
            log.warning("CSV do estoque respondeu HTTP %s", d.status_code)
            return {}
        RAW.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(d.content)
    linhas = list(csv.DictReader(io.StringIO(destino.read_text(encoding="utf-8-sig")), delimiter=";"))
    if not linhas:
        return {}

    por_mes: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for x in linhas:
        v = _num(x["Valor do Estoque"])
        if v is None:
            continue
        m = x["Mes do Estoque"]
        por_mes[m]["total"] += v
        por_mes[m]["externa" if "Externa" in x["Tipo de Divida"] else "interna"] += v
    meses = sorted(por_mes, key=_ordem_mes)
    ultimo = meses[-1]

    venc: dict[str, float] = defaultdict(float)
    indexador: dict[str, float] = defaultdict(float)
    for x in linhas:
        if x["Mes do Estoque"] != ultimo:
            continue
        v = _num(x["Valor do Estoque"])
        if v is None:
            continue
        venc[x["Vencimento do Titulo/Contrato"][6:]] += v
        if "Externa" in x["Tipo de Divida"]:
            rotulo = "Câmbio (dívida externa)"
        else:
            titulo = x["Titulo/Contrato"].upper()
            rotulo = next((r for p, r in INDEXADOR if titulo.startswith(p)), "Outros")
        indexador[rotulo] += v

    ano_ult, mes_ult = ultimo[3:], ultimo[:2]
    doze = sum(v for a, v in venc.items() if a and int(a) <= int(ano_ult) + 1)
    return {
        "meses": [{"mes": m, **{k: round(v, 2) for k, v in por_mes[m].items()}} for m in meses],
        "ultimo_mes": ultimo, "ultimo_ano": int(ano_ult), "ultimo_mes_numero": int(mes_ult),
        "vencimento_por_ano": {a: round(v, 2) for a, v in sorted(venc.items()) if a},
        "vence_ate_o_fim_do_ano_que_vem": round(doze, 2),
        "por_indexador": dict(sorted(((k, round(v, 2)) for k, v in indexador.items()), key=lambda x: -x[1])),
        "fonte": "Tesouro Transparente — Estoque da Dívida Pública Federal (um registro por título e mês)",
        "url": "https://www.tesourotransparente.gov.br/ckan/dataset/estoque-da-divida-publica-federal",
    }


def run() -> Path:
    PUB.mkdir(parents=True, exist_ok=True)
    with httpx.Client(timeout=600, follow_redirects=True) as c:
        juros = juros_nominais(c)
        est = estoque(c)
    if not juros or not est:
        raise SystemExit("faltou uma das fontes da dívida — veja os avisos acima")

    anos = sorted(juros)
    dados = {
        "gerado_em_utc": datetime.now(timezone.utc).isoformat(timespec="minutes"),
        "juros_nominais": juros, "anos_juros": anos, "estoque": est,
        "fonte_juros": "Tesouro Nacional — Resultado do Tesouro Nacional, Série Histórica, tabela 2.1, linha "
                       "'9. JUROS NOMINAIS' (Governo Central, valores correntes)",
        "url_juros": "https://www.tesourotransparente.gov.br/ckan/dataset/resultado-do-tesouro-nacional",
        "nota": "Juros nominais é o custo econômico do ano e não coincide com o que o Orçamento paga de juros "
                "(parte é apropriada no estoque dos títulos). Amortização não é custo: é rolagem.",
    }
    out = PUB / "divida.json"
    out.write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
    log.info("dívida: juros nominais %s–%s (%s = R$ %.1f bi); estoque %s = R$ %.2f tri",
             anos[0], anos[-1], anos[-1], juros[anos[-1]] / 1e9, est["ultimo_mes"],
             est["meses"][-1]["total"] / 1e12)
    return out


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)
    run()
