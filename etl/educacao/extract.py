"""
Educação — IDEB por estado e município (INEP) e escolaridade da população adulta (IBGE, PNAD Contínua).

    python -m etl.educacao.extract [--ano 2023]      # gera data/public/educacao.json

Fonte: INEP, planilhas de divulgação do IDEB (download.inep.gov.br/ideb/resultados/):
- divulgacao_brasil_ideb_<ano>.zip e divulgacao_regioes_ufs_ideb_<ano>.zip — 3 abas (anos iniciais, anos finais, ensino médio), por região/UF e rede
  (Total, Pública, Privada, Estadual, Municipal, Federal), com o IDEB observado de todas as edições (2005 → <ano>).
- divulgacao_{anos_iniciais,anos_finais,ensino_medio}_municipios_<ano>.zip — 1 linha por município × rede.
Cada zip traz o md5 do xlsx; conferimos.

IDEB = fluxo (taxa de aprovação) × desempenho (Saeb, português e matemática), escala 0–10. Bienal; o INEP divulga ~1 ano depois.
Município: usamos a rede PÚBLICA (estadual + municipal); escolas privadas não entram no IDEB municipal.
TLS: o certificado de download.inep.gov.br não fecha a cadeia em algumas máquinas; tentamos verificar e, se falhar, baixamos sem
verificação e conferimos o md5 publicado dentro do zip.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import httpx
import openpyxl

from etl.ufs import NOME_PARA_UF

log = logging.getLogger("etl.educacao")
PUB = Path("data/public")
BASE = "https://download.inep.gov.br/ideb/resultados/"
UA = {"user-agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36"}
ETAPAS = {"anos_iniciais": "ai", "anos_finais": "af", "ensino_medio": "em"}
ABREV = {"M. G. do Sul": "Mato Grosso do Sul", "R. G. do Norte": "Rio Grande do Norte", "R. G. do Sul": "Rio Grande do Sul"}  # como o INEP abrevia
REDES = {"Total": "total", "Pública": "publica", "Privada": "privada", "Estadual": "estadual", "Municipal": "municipal", "Federal": "federal"}


def baixar(c: httpx.Client, nome: str) -> bytes:
    url = BASE + nome
    try:
        r = c.get(url); r.raise_for_status(); return r.content
    except httpx.ConnectError as e:
        if "CERTIFICATE_VERIFY_FAILED" not in str(e):
            raise
        log.warning("TLS do INEP não verificou (%s); baixando sem verificação e conferindo md5", nome)
        with httpx.Client(headers=UA, timeout=600, follow_redirects=True, verify=False) as c2:
            r = c2.get(url); r.raise_for_status(); return r.content


def xlsx_do_zip(conteudo: bytes) -> bytes:
    z = zipfile.ZipFile(io.BytesIO(conteudo))
    xlsx = next(n for n in z.namelist() if n.endswith(".xlsx"))
    md5 = next((n for n in z.namelist() if n.startswith(("md5", xlsx.rsplit("/", 1)[0] + "/md5"))), None)
    dados = z.read(xlsx)
    if md5:
        esperado = re.search(r"[0-9a-f]{32}", z.read(md5).decode("utf-8", "ignore"))
        if esperado and hashlib.md5(dados).hexdigest() != esperado.group(0):
            raise RuntimeError(f"md5 divergente em {xlsx}")
    return dados


def _rede(s) -> str | None:
    return REDES.get(re.sub(r"\s*\(\d\)", "", str(s or "")).strip())


def _val(v):
    try:
        return round(float(v), 1)
    except (TypeError, ValueError):
        return None


def ler_ufs(xlsx: bytes) -> tuple[dict, dict, dict]:
    """Devolve (estados, regioes, brasil): uf → {etapa: {rede: {ano: ideb}}}."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    estados: dict[str, dict] = {}; regioes: dict[str, dict] = {}; brasil: dict = {}
    for ws, etapa in zip(wb.worksheets, ETAPAS.values()):
        rows = ws.iter_rows(values_only=True)
        codigos = None
        for r in rows:
            if r and any(str(x).startswith("VL_OBSERVADO_") for x in r):
                codigos = [str(x) for x in r]; break
        obs = {i: int(m.group(1)) for i, c in enumerate(codigos) if (m := re.fullmatch(r"VL_OBSERVADO_(\d{4})", c))}
        for r in rows:
            if not r or not r[0]:
                continue
            nome, rede = ABREV.get(str(r[0]).strip(), str(r[0]).strip()), _rede(r[1])
            if not rede:
                continue
            serie = {str(a): _val(r[i]) for i, a in obs.items() if i < len(r) and _val(r[i]) is not None}
            if nome == "Brasil":
                alvo = brasil
            elif nome in NOME_PARA_UF:
                alvo = estados.setdefault(NOME_PARA_UF[nome], {"uf": NOME_PARA_UF[nome], "nome": nome})
            elif nome in ("Norte", "Nordeste", "Sudeste", "Sul", "Centro-Oeste"):
                alvo = regioes.setdefault(nome, {"nome": nome})
            else:
                continue
            alvo.setdefault(etapa, {})[rede] = serie
    return estados, regioes, brasil


def ler_municipios(xlsx: bytes, etapa: str, ano: int, acc: dict) -> None:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    codigos = None
    for r in rows:
        if r and "SG_UF" in [str(x) for x in r]:
            codigos = [str(x) for x in r]; break
    i_uf, i_cod, i_nome, i_rede = (codigos.index(k) for k in ("SG_UF", "CO_MUNICIPIO", "NO_MUNICIPIO", "REDE"))
    i_obs = codigos.index(f"VL_OBSERVADO_{ano}")
    for r in rows:
        if not r or _rede(r[i_rede]) != "publica" or len(r) <= i_obs:
            continue
        v = _val(r[i_obs])
        if v is None:
            continue
        cod = int(float(r[i_cod]))
        m = acc.setdefault(cod, {"cod_ibge": cod, "nome": str(r[i_nome]).strip(), "uf": str(r[i_uf]).strip()})
        m[etapa] = v



# --- Escolaridade e analfabetismo (IBGE, PNAD Contínua anual, módulo Educação) ---------------------------------
# O IDEB mede a escola; estas duas tabelas medem a população adulta. São do mesmo IBGE do painel Renda, via SIDRA.
#   7189 — pessoas de 25 anos ou mais por nível de instrução (mil pessoas e distribuição %), 2016 → , Brasil e UF.
#   7113 — taxa de analfabetismo das pessoas de 15 anos ou mais (%), 2016 → , Brasil e UF.
# Ids das variáveis e categorias fixos aqui porque a API de metadados do IBGE dá timeout nos runners do GitHub
# (mesma solução já usada em Renda, Saúde e Inflação). Metadados:
#   https://servicodados.ibge.gov.br/api/v3/agregados/7189/metadados
#   https://servicodados.ibge.gov.br/api/v3/agregados/7113/metadados
SIDRA = "https://apisidra.ibge.gov.br/values/t/{t}/{n}/all/v/{v}/p/{p}{c}?formato=json"
INSTRUCAO = {           # código da classificação 1568 → chave e rótulo
    18837: ("sem_instrucao", "Sem instrução"),
    11779: ("fundamental_incompleto", "Fundamental incompleto"),
    11628: ("fundamental_completo", "Fundamental completo"),
    11629: ("medio_incompleto", "Médio incompleto"),
    11630: ("medio_completo", "Médio completo"),
    11631: ("superior_incompleto", "Superior incompleto"),
    99713: ("superior_completo", "Superior completo"),
}
ORDEM_INSTRUCAO = ["sem_instrucao", "fundamental_incompleto", "fundamental_completo", "medio_incompleto",
                   "medio_completo", "superior_incompleto", "superior_completo"]


def _sidra(c: httpx.Client, t: int, n: str, v: int, p: str, c_: str = "") -> list[dict]:
    erro = None
    for tentativa in range(3):
        try:
            r = c.get(SIDRA.format(t=t, n=n, v=v, p=p, c=c_)); r.raise_for_status()
            return r.json()[1:]
        except (httpx.TransportError, httpx.HTTPStatusError) as e:
            erro = e; log.warning("SIDRA t%s tentativa %d: %s", t, tentativa + 1, e)
    raise erro


def _num(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def escolaridade(c: httpx.Client) -> dict:
    """Até onde a população adulta estudou e quantos não sabem ler, por ano e por estado."""
    from etl.ufs import NOME_PARA_UF

    # 1) nível de instrução dos 25 anos ou mais: série do Brasil (% e mil pessoas) e último ano por UF
    serie: dict[str, dict[str, float]] = {}
    pessoas: dict[str, dict[str, float]] = {}
    for cod, (chave, _) in INSTRUCAO.items():
        for var, destino in ((10270, serie), (1643, pessoas)):
            for linha in _sidra(c, 7189, "n1", var, "all", f"/c1568/{cod}/c2/6794"):
                valor = _num(linha.get("V"))
                if valor is not None:
                    destino.setdefault(chave, {})[linha["D3N"]] = valor
    anos = sorted({a for v in serie.values() for a in v})
    ultimo = anos[-1] if anos else None

    por_uf: dict[str, dict] = {}
    for cod, (chave, _) in INSTRUCAO.items():
        for linha in _sidra(c, 7189, "n3", 10270, "last 1", f"/c1568/{cod}/c2/6794"):
            uf = NOME_PARA_UF.get((linha.get("D1N") or "").strip())
            valor = _num(linha.get("V"))
            if uf and valor is not None:
                por_uf.setdefault(uf, {"uf": uf, "nome": linha["D1N"].strip()})[chave] = valor

    # 2) taxa de analfabetismo das pessoas de 15 anos ou mais
    analf_serie = {l["D3N"]: _num(l.get("V")) for l in _sidra(c, 7113, "n1", 10267, "all", "/c2/6794/c58/2795")
                   if _num(l.get("V")) is not None}
    for linha in _sidra(c, 7113, "n3", 10267, "last 1", "/c2/6794/c58/2795"):
        uf = NOME_PARA_UF.get((linha.get("D1N") or "").strip())
        valor = _num(linha.get("V"))
        if uf and valor is not None:
            por_uf.setdefault(uf, {"uf": uf, "nome": linha["D1N"].strip()})["analfabetismo"] = valor

    # 3) o que a série diz do começo ao fim (a página mostra, não conclui)
    def _var(d: dict) -> dict | None:
        if not d or len(d) < 2:
            return None
        a0, a1 = min(d), max(d)
        return {"de": a0, "para": a1, "valor_de": d[a0], "valor_para": d[a1], "diferenca": round(d[a1] - d[a0], 2)}

    mudanca = {k: _var(v) for k, v in serie.items()}
    mudanca["analfabetismo"] = _var(analf_serie)
    sem_fundamental = {a: round(sum(serie.get(k, {}).get(a, 0) for k in ("sem_instrucao", "fundamental_incompleto")), 2)
                       for a in anos}
    log.info("escolaridade %s: %.1f%% dos adultos com médio completo ou mais, %.1f%% sem o fundamental; analfabetismo %.1f%%",
             ultimo,
             sum(serie.get(k, {}).get(ultimo, 0) for k in ("medio_completo", "superior_incompleto", "superior_completo")),
             sem_fundamental.get(ultimo, 0), analf_serie.get(ultimo, 0))
    return {
        "ano": ultimo, "anos": anos, "ordem": ORDEM_INSTRUCAO,
        "rotulos": {k: r for k, r in INSTRUCAO.values()},
        "distribuicao": serie, "pessoas_mil": pessoas,
        "sem_fundamental": sem_fundamental,
        "analfabetismo": analf_serie, "mudanca": mudanca,
        "por_uf": sorted(por_uf.values(), key=lambda d: d["uf"]),
        "fonte_longo": "IBGE: Censo Demográfico (tabela 204, 1991 e 2000), PNAD antiga (tabela 271, 2001–2015) e PNAD "
                       "Contínua anual (tabela 7113, 2016 em diante). Taxa de 15 anos ou mais somada das faixas de idade "
                       "nas duas primeiras, que não publicam o agregado.",
        "fonte": "IBGE, PNAD Contínua anual (módulo Educação): tabela 7189 (pessoas de 25 anos ou mais por nível de "
                 "instrução) e 7113 (taxa de analfabetismo das pessoas de 15 anos ou mais)",
    }



# --- Analfabetismo desde 1991, juntando três pesquisas do IBGE ----------------------------------------------------
# Não existe uma série única: o IBGE mediu analfabetismo no Censo, na PNAD antiga e na PNAD Contínua, com metodologias
# diferentes. Montamos os três trechos e a página marca onde a pesquisa troca — é o mais longe que o dado oficial vai,
# e 1991 é o primeiro Censo depois da Constituição de 1988.
#   204 (Censo 1991 e 2000) — população de 5 anos ou mais por alfabetização e grupo de idade
#   271 (PNAD 2001–2015)    — pessoas de 5 anos ou mais por alfabetização e grupo de idade
#   7113 (PNAD Contínua 2016–) — taxa de analfabetismo já calculada
# A taxa de 15 anos ou mais é somada por nós a partir das faixas de idade (as duas primeiras tabelas não publicam o
# agregado "15 anos ou mais"): não alfabetizadas ÷ total, nas faixas de 15 anos para cima.
FAIXAS_15_MAIS = {
    204: "1143,1144,1145,1146,1147,1148,1149,1150,1151,1152,1153,1154,1155,2503",
    271: "1143,1144,1145,3299,3300,3301,3302",
}
TOTAL_ALFAB = {204: 0, 271: 93024}          # código de "Total" na classificação 59, que muda entre as duas tabelas


def analfabetismo_longo(c: httpx.Client) -> dict:
    """Taxa de analfabetismo das pessoas de 15 anos ou mais em cada pesquisa que o IBGE já usou, de 1991 para cá."""
    trechos: dict[str, dict] = {}
    for tabela, variavel, rotulo in ((204, 157, "Censo Demográfico"), (271, 121, "PNAD (antiga)")):
        soma: dict[str, dict[str, float]] = {}
        for cod, chave in ((1024, "nao_alfabetizadas"), (TOTAL_ALFAB[tabela], "total")):
            extra = "/c1/0/c2/0" if tabela == 271 else ""
            for linha in _sidra(c, tabela, "n1", variavel, "all", f"/c59/{cod}/c58/{FAIXAS_15_MAIS[tabela]}{extra}"):
                valor = _num(linha.get("V"))
                if valor is not None:
                    ano = linha["D3N"]          # D3N é o Ano nas duas tabelas (D2N é o nome da variável)
                    soma.setdefault(ano, {}).setdefault(chave, 0.0)
                    soma[ano][chave] += valor
        for ano, v in soma.items():
            if v.get("total"):
                trechos.setdefault(rotulo, {})[ano] = round(100 * v["nao_alfabetizadas"] / v["total"], 2)
    log.info("analfabetismo longo: %s", {k: f"{min(v)}–{max(v)}" for k, v in trechos.items()})
    return {k: dict(sorted(v.items())) for k, v in trechos.items()}


def run(ano: int) -> Path:
    PUB.mkdir(parents=True, exist_ok=True)
    with httpx.Client(headers=UA, timeout=600, follow_redirects=True) as c:
        estados, regioes, _ = ler_ufs(xlsx_do_zip(baixar(c, f"divulgacao_regioes_ufs_ideb_{ano}.zip")))
        _, _, brasil = ler_ufs(xlsx_do_zip(baixar(c, f"divulgacao_brasil_ideb_{ano}.zip")))
        municipios: dict[int, dict] = {}
        for nome, etapa in ETAPAS.items():
            ler_municipios(xlsx_do_zip(baixar(c, f"divulgacao_{nome}_municipios_{ano}.zip")), etapa, ano, municipios)
            log.info("%s: %d municípios com IDEB público", nome, sum(1 for m in municipios.values() if etapa in m))
        esc = escolaridade(c)
        esc["analfabetismo_longo"] = analfabetismo_longo(c)
    dados = {"gerado_em_utc": datetime.now(timezone.utc).isoformat(timespec="minutes"), "ano": ano,
             "brasil": brasil, "regioes": list(regioes.values()), "estados": sorted(estados.values(), key=lambda e: e["uf"]),
             "municipios": sorted(municipios.values(), key=lambda m: m["cod_ibge"]),
             "escolaridade": esc,
             "fontes": {"inep": f"INEP, IDEB {ano}, planilhas de divulgação por região/UF e por município (rede pública)",
                        "ibge": esc["fonte"], "url": BASE}}
    out = PUB / "educacao.json"
    out.write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
    log.info("IDEB %d: Brasil AI %s, AF %s, EM %s; %d UFs; %d municípios", ano, brasil.get("ai", {}).get("total", {}).get(str(ano)),
             brasil.get("af", {}).get("total", {}).get(str(ano)), brasil.get("em", {}).get("total", {}).get(str(ano)), len(estados), len(municipios))
    return out


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)
    ap = argparse.ArgumentParser(); ap.add_argument("--ano", type=int, default=2023)
    run(ap.parse_args().ano)
