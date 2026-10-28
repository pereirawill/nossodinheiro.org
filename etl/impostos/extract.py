"""
Impostos — arrecadação federal por tributo, base de incidência e estado (Receita Federal, dados abertos, sem chave).

    python -m etl.impostos.extract            # baixa o CSV e gera data/public/impostos.json

Fonte: Receita Federal, "Arrecadação por estado" — https://www.gov.br/receitafederal/dados/arrecadacao-estado.csv
(latin-1, ';', valores em R$; Ano, Mês, UF e uma coluna por tributo; mensal desde 2000, atualizado mensalmente).
É a arrecadação administrada pela RFB (impostos e contribuições federais + previdência). NÃO inclui tributos de
estados e municípios (ICMS, IPVA, ISS, IPTU) — ver pendências na metodologia.

Agrupamento por BASE DE INCIDÊNCIA (classificação econômica usual; quem recolhe ≠ quem paga, e a página avisa):
  renda_trabalho   IRPF; IRRF sobre rendimentos do trabalho; contribuição dos servidores (CPSSS)
  folha            Receita previdenciária (INSS: parte do empregado + parte da empresa; a fonte não separa)
  lucro_empresas   IRPJ; CSLL
  consumo          IPI; COFINS; PIS/PASEP; CIDE-combustíveis; impostos de importação e exportação
  renda_capital    IRRF sobre rendimentos do capital, remessas ao exterior e outros rendimentos; IOF
  patrimonio       ITR
  simples_outros   Pagamento unificado (Simples Nacional), retenção na fonte Lei 10.833, REFIS/PAES, FUNDAF,
                   outras receitas administradas, demais receitas, administradas por outros órgãos, CPMF/IPMF (extintos)

Saída: data/public/impostos.json — {gerado_em_utc, ultimo_mes:{ano, mes}, anual:[{ano, meses, total, bases:{…}, tributos:{…}}],
  mensal:[{ano, mes, total, bases:{…}}] (últimos 24 meses), por_uf:[{uf, total, por_habitante?}] (ano corrente),
  tributos_ultimo_ano:[{tributo, base, valor}], carga:{serie, por_ente, bases:{linhas:[{base, federal, estadual,
  municipal, total_milhoes, pct}]}}}
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import httpx

CTB = "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/estudos/carga-tributaria/tabelas-carga-tributaria-no-brasil-{ano}/@@download/file"

log = logging.getLogger("etl.impostos")
PUB = Path("data/public")
URL = "https://www.gov.br/receitafederal/dados/arrecadacao-estado.csv"
HEADERS = {"user-agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36"}
MESES = {m: i for i, m in enumerate(["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"], 1)}

BASES = {
    "renda_trabalho": ["IRPF", "IRRF - RENDIMENTOS DO TRABALHO", "CONTRIBUIÇÃO PLANO SEG. SOC. SERVIDORES", "CPSSS - Contrib. p/ o Plano de Segurid. Social Serv. Público"],
    "folha": ["RECEITA PREVIDENCIÁRIA", "RECEITA PREVIDENCIÁRIA - PRÓPRIA", "RECEITA PREVIDENCIÁRIA - DEMAIS"],
    "lucro_empresas": ["IRPJ - ENTIDADES FINANCEIRAS", "IRPJ - DEMAIS EMPRESAS", "CSLL", "CSLL - FINANCEIRAS", "CSLL - DEMAIS"],
    "consumo": ["IPI - FUMO", "IPI - BEBIDAS", "IPI - AUTOMÓVEIS", "IPI - VINCULADO À IMPORTACAO", "IPI - OUTROS", "COFINS", "COFINS - FINANCEIRAS", "COFINS - DEMAIS",
                "CONTRIBUIÇÃO PARA O PIS/PASEP", "CONTRIBUIÇÃO PARA O PIS/PASEP - FINANCEIRAS", "CONTRIBUIÇÃO PARA O PIS/PASEP - DEMAIS",
                "CIDE-COMBUSTÍVEIS (parc. não dedutível)", "CIDE-COMBUSTÍVEIS", "IMPOSTO SOBRE IMPORTAÇÃO", "IMPOSTO SOBRE EXPORTAÇÃO"],
    "renda_capital": ["IRRF - RENDIMENTOS DO CAPITAL", "IRRF - REMESSAS P/ EXTERIOR", "IRRF - OUTROS RENDIMENTOS", "IMPOSTO S/ OPERAÇÕES FINANCEIRAS"],
    "patrimonio": ["IMPOSTO TERRITORIAL RURAL"],
    "simples_outros": ["PAGAMENTO UNIFICADO", "RETENÇÃO NA FONTE - LEI 10.833, Art. 30", "REFIS", "PAES", "CONTRIBUICÕES PARA FUNDAF",
                       "OUTRAS RECEITAS ADMINISTRADAS", "DEMAIS RECEITAS", "ADMINISTRADAS POR OUTROS ÓRGÃOS", "CPMF", "IMPOSTO PROVIS.S/ MOVIMENT. FINANC. - IPMF"],
}
BASE_DE = {t: b for b, ts in BASES.items() for t in ts}
ROTULO_BASE = {"renda_trabalho": "Renda do trabalho", "folha": "Folha (previdência)", "lucro_empresas": "Lucro das empresas", "consumo": "Consumo e produção",
               "renda_capital": "Renda do capital e operações financeiras", "patrimonio": "Patrimônio", "simples_outros": "Simples Nacional e outros"}
# Nomes curtos para a tabela de tributos
NOME_TRIBUTO = {"IRPF": "IRPF (pessoa física)", "IRRF - RENDIMENTOS DO TRABALHO": "IRRF sobre salários", "IRPJ - DEMAIS EMPRESAS": "IRPJ (empresas)",
                "IRPJ - ENTIDADES FINANCEIRAS": "IRPJ (bancos)", "CSLL - DEMAIS": "CSLL (empresas)", "CSLL - FINANCEIRAS": "CSLL (bancos)",
                "COFINS - DEMAIS": "COFINS", "COFINS - FINANCEIRAS": "COFINS (bancos)", "CONTRIBUIÇÃO PARA O PIS/PASEP - DEMAIS": "PIS/PASEP",
                "CONTRIBUIÇÃO PARA O PIS/PASEP - FINANCEIRAS": "PIS/PASEP (bancos)", "IMPOSTO S/ OPERAÇÕES FINANCEIRAS": "IOF",
                "RECEITA PREVIDENCIÁRIA": "Contribuição previdenciária (INSS)", "CPSSS - Contrib. p/ o Plano de Segurid. Social Serv. Público": "Previdência dos servidores (CPSSS)",
                "PAGAMENTO UNIFICADO": "Simples Nacional (pagamento unificado)", "IMPOSTO SOBRE IMPORTAÇÃO": "Imposto de importação", "IMPOSTO TERRITORIAL RURAL": "ITR",
                "IRRF - RENDIMENTOS DO CAPITAL": "IRRF sobre rendimentos do capital", "IRRF - REMESSAS P/ EXTERIOR": "IRRF sobre remessas ao exterior",
                "IRRF - OUTROS RENDIMENTOS": "IRRF outros rendimentos", "CIDE-COMBUSTÍVEIS": "CIDE-combustíveis", "IPI - OUTROS": "IPI (outros)",
                "IPI - VINCULADO À IMPORTACAO": "IPI sobre importados", "IPI - AUTOMÓVEIS": "IPI automóveis", "IPI - BEBIDAS": "IPI bebidas", "IPI - FUMO": "IPI fumo",
                "OUTRAS RECEITAS ADMINISTRADAS": "Outras receitas administradas", "ADMINISTRADAS POR OUTROS ÓRGÃOS": "Administradas por outros órgãos",
                "RETENÇÃO NA FONTE - LEI 10.833, Art. 30": "Retenção na fonte (Lei 10.833)", "IMPOSTO SOBRE EXPORTAÇÃO": "Imposto de exportação"}
POPULACAO = 212_600_000  # IBGE, estimativa (mesma do painel Poderes)


def _brl(s: str) -> float:
    """O CSV mistura formatos: '65512,84' e '8762245,567' (decimal com vírgula, 1 a 3 casas), '47.953.915' e
    '1,006,523,961' (milhares sem decimais), '1.234,56'. Regra: com os dois separadores, o último é o decimal; só vírgula:
    uma = decimal, várias = milhar; só ponto: vários = milhar, um seguido de 3 dígitos = milhar, senão decimal."""
    s = (s or "").strip()
    if not s or s == "-":
        return 0.0
    if "." in s and "," in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        mil = "." if dec == "," else ","
        return float(s.replace(mil, "").replace(dec, "."))
    if "," in s:                       # vírgula única = decimal (qualquer nº de casas: '8762245,567'); várias = milhar
        return float(s.replace(",", ".")) if s.count(",") == 1 else float(s.replace(",", ""))
    if "." in s:                       # ponto único seguido de 3 dígitos = milhar ('47.953'); senão decimal; vários = milhar
        partes = s.split(".")
        if len(partes) == 2 and len(partes[1]) != 3:
            return float(s)
        return float(s.replace(".", ""))
    return float(s)


# Famílias com coluna agregada E colunas detalhadas (o CSV preenche uma OU outra conforme o ano): sem isto, soma em dobro.
FAMILIAS = {
    "COFINS": ["COFINS - FINANCEIRAS", "COFINS - DEMAIS"],
    "CONTRIBUIÇÃO PARA O PIS/PASEP": ["CONTRIBUIÇÃO PARA O PIS/PASEP - FINANCEIRAS", "CONTRIBUIÇÃO PARA O PIS/PASEP - DEMAIS"],
    "CSLL": ["CSLL - FINANCEIRAS", "CSLL - DEMAIS"],
    "CIDE-COMBUSTÍVEIS": ["CIDE-COMBUSTÍVEIS (parc. não dedutível)"],
    "CONTRIBUIÇÃO PLANO SEG. SOC. SERVIDORES": ["CPSSS - Contrib. p/ o Plano de Segurid. Social Serv. Público"],
    "RECEITA PREVIDENCIÁRIA": ["RECEITA PREVIDENCIÁRIA - PRÓPRIA", "RECEITA PREVIDENCIÁRIA - DEMAIS"],
}


def _sem_duplicidade(valores: dict[str, float], brutos: dict[str, str]) -> dict[str, float]:
    """Se a coluna agregada veio preenchida, zera as detalhadas; senão, mantém as detalhadas. Nunca as duas."""
    v = dict(valores)
    for agg, subs in FAMILIAS.items():
        if agg in v and (brutos.get(agg) or "").strip() not in ("", "-"):
            for s in subs:
                v[s] = 0.0
        else:
            v[agg] = 0.0
    return v


def carregar(txt: str) -> list[dict]:
    rows = list(csv.DictReader(io.StringIO(txt), delimiter=";"))
    tributos = [k for k in rows[0].keys() if k not in ("Ano", "Mês", "UF")]
    desconhecidos = [t for t in tributos if t not in BASE_DE]
    if desconhecidos:
        log.warning("tributos sem base definida (vão para 'simples_outros'): %s", desconhecidos)
    out = []
    for r in rows:
        valores = _sem_duplicidade({t: _brl(r[t]) for t in tributos}, r)
        out.append({"ano": int(r["Ano"]), "mes": MESES[r["Mês"].strip()], "uf": r["UF"].strip(), "valores": valores})
    return out


def agregar(linhas: list[dict]) -> dict:
    anual: dict[int, dict] = {}
    mensal: dict[tuple[int, int], dict] = {}
    for l in linhas:
        a = anual.setdefault(l["ano"], {"ano": l["ano"], "meses": set(), "total": 0.0, "bases": defaultdict(float), "tributos": defaultdict(float)})
        m = mensal.setdefault((l["ano"], l["mes"]), {"ano": l["ano"], "mes": l["mes"], "total": 0.0, "bases": defaultdict(float)})
        a["meses"].add(l["mes"])
        for t, v in l["valores"].items():
            b = BASE_DE.get(t, "simples_outros")
            a["total"] += v; a["bases"][b] += v; a["tributos"][t] += v
            m["total"] += v; m["bases"][b] += v
    anos = sorted(anual)
    ultimo_ano = anos[-1]; ultimo_mes = max(anual[ultimo_ano]["meses"])
    anual_out = [{"ano": a, "meses": len(anual[a]["meses"]), "total": round(anual[a]["total"], 2),
                  "bases": {b: round(v, 2) for b, v in anual[a]["bases"].items()},
                  "tributos": {t: round(v, 2) for t, v in anual[a]["tributos"].items() if v}} for a in anos]
    chaves = sorted(mensal)[-24:]
    mensal_out = [{"ano": k[0], "mes": k[1], "total": round(mensal[k]["total"], 2), "bases": {b: round(v, 2) for b, v in mensal[k]["bases"].items()}} for k in chaves]
    por_uf: dict[str, float] = defaultdict(float)
    for l in linhas:
        if l["ano"] == ultimo_ano:
            por_uf[l["uf"]] += sum(l["valores"].values())
    por_uf_out = sorted([{"uf": uf, "total": round(v, 2)} for uf, v in por_uf.items()], key=lambda x: -x["total"])
    trib = anual[ultimo_ano]["tributos"]
    tributos_out = sorted([{"tributo": NOME_TRIBUTO.get(t, t.title()), "codigo": t, "base": BASE_DE.get(t, "simples_outros"), "valor": round(v, 2)} for t, v in trib.items() if v > 0], key=lambda x: -x["valor"])
    return {"gerado_em_utc": datetime.now(timezone.utc).isoformat(timespec="minutes"), "fonte": "Receita Federal, Arrecadação por estado (dados abertos)", "url": URL,
            "populacao": POPULACAO, "rotulo_base": ROTULO_BASE, "ultimo_mes": {"ano": ultimo_ano, "mes": ultimo_mes},
            "anual": anual_out, "mensal": mensal_out, "por_uf": por_uf_out, "tributos_ultimo_ano": tributos_out}


def _texto(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


# Bases de incidência da tabela INC 01-A da planilha de carga tributária: a única divisão OFICIAL que cobre os três
# níveis de governo (a classificação por família de tributo no topo deste arquivo é nossa e só cobre a receita federal).
BASES_OFICIAIS = {"1000": "Renda", "2000": "Folha de salários", "3000": "Propriedade",
                  "4000": "Bens e serviços", "5000": "Transações financeiras", "9000": "Outros tributos"}


def bases_de_incidencia(wb) -> dict | None:
    """
    Receita tributária por base de incidência e nível de governo (tabela INC 01-A, em R$ milhões). Código da base na
    coluna 1 (nível 1) ou 2 (detalhe); rótulo na 3; valores federal/estadual/municipal/total nas colunas 4 a 7.
    Células "n/a" = o nível não arrecada aquele tributo.
    """
    if "INC01A" not in wb.sheetnames:
        log.warning("planilha de carga tributária sem a aba INC01A")
        return None
    def _cod(v) -> str:                                  # o código vem como inteiro (0, 1000, 2000…)
        return str(int(v)) if isinstance(v, (int, float)) else _texto(v)

    linhas, total = [], None
    for l in wb["INC01A"].iter_rows(max_col=8, values_only=True):
        cod = _cod(l[1] if len(l) > 1 else "")
        if cod == "0":
            total = float(l[7]) if isinstance(l[7], (int, float)) else None
        elif cod in BASES_OFICIAIS and isinstance(l[7], (int, float)):
            linhas.append({"base": BASES_OFICIAIS[cod], "codigo": cod,
                           "federal": round(float(l[4]), 2) if isinstance(l[4], (int, float)) else None,
                           "estadual": round(float(l[5]), 2) if isinstance(l[5], (int, float)) else None,
                           "municipal": round(float(l[6]), 2) if isinstance(l[6], (int, float)) else None,
                           "total_milhoes": round(float(l[7]), 2)})
    if not linhas or not total:
        log.warning("INC01A: nenhuma base de incidência lida")
        return None
    soma = sum(i["total_milhoes"] for i in linhas)
    if abs(soma - total) > total * 0.005:
        raise SystemExit(f"bases de incidência somam {soma:.0f} contra total {total:.0f} (INC01A)")
    for i in linhas:
        i["pct"] = round(i["total_milhoes"] / total, 4)
    log.info("bases de incidência (INC01A): %s", ", ".join(f"{i['base']} {100*i['pct']:.1f}%" for i in linhas))
    return {"total_milhoes": round(total, 2), "tabela": "INC 01-A", "linhas": linhas}


def serie_longa(wb) -> dict[str, float]:
    """
    Carga tributária ano a ano, da tabela TRIB 01-B (receita por tributo em % do PIB): a linha "Total da Receita
    Tributária" dá 2002 → último ano em uma única tabela, sem emendar edições do estudo. A Tab_01 traz só os anos
    recentes e alguns de referência, e é por isso que a série dela pula anos.
    """
    if "T01B" not in wb.sheetnames:
        return {}
    linhas = [list(r) for r in wb["T01B"].iter_rows(max_col=40, values_only=True)]
    cab = next((l for l in linhas if sum(1 for v in l if _texto(v).isdigit() and len(_texto(v)) == 4) >= 5), None)
    if not cab:
        log.warning("T01B: linha de anos não encontrada")
        return {}
    anos = {i: _texto(v) for i, v in enumerate(cab) if _texto(v).isdigit() and len(_texto(v)) == 4}
    # o rótulo fica na coluna 2 ou 3 conforme a aba; procuramos nas primeiras colunas
    total = next((l for l in linhas
                  if any(_texto(v).startswith("Total da Receita Tribut") for v in l[:4])), None)
    if not total:
        log.warning("T01B: linha de total não encontrada")
        return {}
    saida = {ano: round(float(total[i]), 6) for i, ano in anos.items()
             if i < len(total) and isinstance(total[i], (int, float))}
    log.info("série longa da carga (T01B): %d anos, %s a %s", len(saida), min(saida), max(saida))
    return saida


def ocde(wb) -> dict | None:
    """Brasil × média da OCDE por base de incidência (Tab_07, que a Receita monta com o Revenue Statistics da OCDE)."""
    if "Tab_07" not in wb.sheetnames:
        return None
    linhas = [list(r) for r in wb["Tab_07"].iter_rows(max_row=12, max_col=12, values_only=True)]
    grupos = next((l for l in linhas if any(_texto(v).startswith("Carga Tribut") for v in l[:4])), None)
    anos_l = next((l for l in linhas if sum(1 for v in l if _texto(v).isdigit() and len(_texto(v)) == 4) >= 4), None)
    if not grupos or not anos_l:
        return None
    # os rótulos de grupo vêm em células mescladas: repete o último rótulo visto
    rotulos, atual = {}, None
    for i, v in enumerate(grupos):
        if _texto(v):
            atual = _texto(v)
        if atual:
            rotulos[i] = atual
    dados: dict[str, dict] = {}
    for l in linhas:
        quem = next((_texto(v) for v in l[:3] if _texto(v)), "")
        if not quem.startswith(("Brasil", "Média OCDE")):
            continue
        chave = "brasil" if quem.startswith("Brasil") else "ocde"
        for i, ano in ((i, _texto(v)) for i, v in enumerate(anos_l) if _texto(v).isdigit()):
            if i < len(l) and isinstance(l[i], (int, float)):
                dados.setdefault(rotulos.get(i, "?"), {}).setdefault(ano, {})[chave] = round(float(l[i]), 2)
    if not dados:
        return None
    log.info("Brasil × OCDE (Tab_07): %d bases, anos %s", len(dados),
             sorted({a for v in dados.values() for a in v}))
    return {"bases": dados, "fonte": "Receita Federal, Carga Tributária no Brasil, tabela 07, a partir do Revenue "
                                     "Statistics da OCDE (média de 36 países)"}


def carga_tributaria(c: httpx.Client) -> dict | None:
    """
    Carga tributária bruta (tudo que União, estados e municípios arrecadam em tributos, sobre o PIB), da planilha oficial
    "Carga Tributária no Brasil" da Receita Federal. Tab_01 traz a série (PIB, arrecadação e carga); Tab_02, a divisão por
    ente federativo no último ano. Os percentuais vêm como fração (0,3220 = 32,20% do PIB).
    """
    import openpyxl

    conteudo = None
    for ano in range(datetime.now().year, datetime.now().year - 4, -1):
        r = c.get(CTB.format(ano=ano))
        if r.status_code == 200 and r.content[:2] == b"PK":
            conteudo, edicao = r.content, ano
            break
    if not conteudo:
        log.warning("planilha de carga tributária não encontrada na Receita Federal")
        return None
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)

    serie: dict[str, dict] = {}
    linhas = [list(r) for r in wb["Tab_01"].iter_rows(values_only=True)]
    cab = next((l for l in linhas if any(_texto(v).isdigit() and len(_texto(v)) == 4 for v in l)), None)
    anos = {i: _texto(v) for i, v in enumerate(cab or []) if _texto(v).isdigit() and len(_texto(v)) == 4}
    rotulos = {"Produto Interno Bruto": "pib", "Arrecadação Tributária Bruta": "arrecadacao", "Carga Tributária Bruta": "carga"}
    for l in linhas:
        chave = next((v for k, v in rotulos.items() if _texto(l[1] if len(l) > 1 else "").startswith(k)), None)
        if not chave:
            continue
        for i, ano in anos.items():
            if i < len(l) and isinstance(l[i], (int, float)):
                serie.setdefault(ano, {})[chave] = round(float(l[i]), 6)

    entes, ano_entes = [], None
    linhas2 = [list(r) for r in wb["Tab_02"].iter_rows(values_only=True)]
    cab2 = next((l for l in linhas2 if any(_texto(v).isdigit() and len(_texto(v)) == 4 for v in l)), [])
    cols_ano = sorted((i, _texto(v)) for i, v in enumerate(cab2) if _texto(v).isdigit() and len(_texto(v)) == 4)
    if cols_ano:
        col, ano_entes = cols_ano[-1]                       # bloco do ano mais recente: arrecadação, % do PIB, % da arrecadação
        for l in linhas2:
            nome = _texto(l[1] if len(l) > 1 else "")
            if nome in ("União", "Estados", "Municípios", "Total") and len(l) > col + 1 and isinstance(l[col], (int, float)):
                entes.append({"ente": nome, "arrecadacao_milhoes": round(float(l[col]), 2),
                              "pct_pib": round(float(l[col + 1]), 6) if isinstance(l[col + 1], (int, float)) else None})
    ultimo = ano_entes or (max(serie) if serie else None)
    log.info("carga tributária (CTB %d): %s = %.2f%% do PIB; %d anos na série; %d entes", edicao, ultimo,
             100 * (next((e["pct_pib"] for e in entes if e["ente"] == "Total"), 0) or serie.get(ultimo, {}).get("carga", 0)), len(serie), len(entes))
    return {"edicao": edicao, "ano": ultimo, "serie": serie, "por_ente": entes,
            "bases": bases_de_incidencia(wb), "serie_longa": serie_longa(wb), "ocde": ocde(wb),
            "fonte": f"Receita Federal, Carga Tributária no Brasil {edicao} (tabelas 01, 02 e INC 01-A)",
            "url": CTB.format(ano=edicao)}


def run() -> Path:
    PUB.mkdir(parents=True, exist_ok=True)
    with httpx.Client(headers=HEADERS, timeout=300, follow_redirects=True) as c:
        r = c.get(URL); r.raise_for_status()
        linhas = carregar(r.content.decode("latin-1"))
        dados = agregar(linhas)
        dados["carga"] = carga_tributaria(c)
    u = dados["ultimo_mes"]; a = dados["anual"][-1]
    if a["total"] < 1e11:
        raise SystemExit(f"total de {u['ano']} = {a['total']:.0f}: unidade/arquivo mudou?")
    out = PUB / "impostos.json"
    out.write_text(json.dumps(dados, ensure_ascii=False), encoding="utf-8")
    log.info("%d linhas; último mês %d-%02d; %d: R$ %.1f bi em %d meses (%s)", len(linhas), u["ano"], u["mes"], a["ano"], a["total"] / 1e9, a["meses"],
             ", ".join(f"{ROTULO_BASE[b]} {v/1e9:.0f} bi" for b, v in sorted(a["bases"].items(), key=lambda kv: -kv[1])))
    return out


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    run()
