"""
Dividendos e participações que a União recebe das estatais, desde 1997.

Fonte: **Tesouro Nacional, Resultado do Tesouro Nacional (RTN) — Série Histórica**, planilha publicada no CKAN do
Tesouro Transparente (`package_show?id=resultado-do-tesouro-nacional`; o nome do arquivo muda a cada mês, por isso é
resolvido pela API). Aba **2.2** (anual, R$ milhões correntes), linha **1.4.2 Dividendos e Participações**, dentro das
receitas não administradas pela Receita Federal, com abertura por empresa nas linhas 1.4.2.1 a 1.4.2.9.

É **regime de caixa**: o que entrou no Tesouro no ano, e não o lucro do ano — dividendo é pago no ano seguinte ao
resultado, e às vezes em parcelas. Não é imposto: é a parte do lucro que cabe ao acionista União.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path

import httpx
import openpyxl

log = logging.getLogger("etl.estatais.dividendos")
RAW = Path("data/raw")
CKAN = ("https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show"
        "?id=resultado-do-tesouro-nacional")
ABA = "2.2"
LINHA_TOTAL = "1.4.2"
RE_CODIGO = re.compile(r"^(1\.4\.2(?:\.\d)?)\s+(.*)$")


def _planilha(c: httpx.Client) -> bytes | None:
    r = c.get(CKAN)
    if r.status_code != 200:
        log.warning("CKAN do Tesouro respondeu HTTP %s", r.status_code)
        return None
    recursos = (r.json().get("result") or {}).get("resources") or []
    url = next((x["url"] for x in recursos if (x.get("format") or "").upper() == "XLSX"
                and "seriehistorica" in (x.get("url") or "").lower()), None)
    if not url:
        log.warning("série histórica do RTN não encontrada no CKAN")
        return None
    destino = RAW / f"rtn_{url.rsplit('/', 1)[-1]}"
    if destino.exists() and destino.stat().st_size > 100_000:
        return destino.read_bytes()
    p = c.get(url)
    if p.status_code != 200:
        log.warning("planilha do RTN respondeu HTTP %s", p.status_code)
        return None
    RAW.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(p.content)
    return p.content


def coletar(c: httpx.Client) -> dict:
    conteudo = _planilha(c)
    if conteudo is None:
        return {}
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    if ABA not in wb.sheetnames:
        log.warning("aba %s não existe na planilha do RTN", ABA)
        return {}
    linhas = list(wb[ABA].iter_rows(values_only=True))
    anos: list[str] = []
    for row in linhas[:12]:
        if row and str(row[0] or "").strip().lower().startswith("discrimina"):
            anos = [str(int(v)) for v in row[1:] if isinstance(v, (int, float)) and 1990 < v < 2100]
            break
    if not anos:
        log.warning("não achei a linha de anos na aba %s", ABA)
        return {}

    total: dict[str, float] = {}
    por_empresa: dict[str, dict[str, float]] = {}
    for row in linhas:
        m = RE_CODIGO.match(str(row[0] or "").strip())
        if not m:
            continue
        codigo, nome = m.group(1), re.sub(r"\s+", " ", m.group(2)).strip(" -")
        serie = {}
        for i, ano in enumerate(anos, start=1):
            v = row[i] if i < len(row) else None
            if isinstance(v, (int, float)):
                serie[ano] = round(float(v) * 1e6, 2)      # a planilha está em R$ milhões
        if codigo == LINHA_TOTAL:
            total = serie
        elif serie:
            por_empresa[nome] = serie

    if not total:
        log.warning("linha %s não encontrada na aba %s", LINHA_TOTAL, ABA)
        return {}
    anos_com_dado = [a for a in anos if a in total]
    ult = anos_com_dado[-1]
    log.info("dividendos da União %s–%s: %s = R$ %.1f bi", anos_com_dado[0], ult, ult, total[ult] / 1e9)
    return {
        "anos": anos_com_dado, "total": total, "por_empresa": por_empresa,
        "fonte": "Tesouro Nacional, Resultado do Tesouro Nacional — Série Histórica, tabela 2.2, linha 1.4.2 "
                 "(Dividendos e Participações), em regime de caixa",
        "url": "https://www.tesourotransparente.gov.br/ckan/dataset/resultado-do-tesouro-nacional",
    }
