"""IRPF: leitura da planilha dos Grandes Números com uma pasta sintética (sem rede)."""
import io

import httpx
import openpyxl
import pytest

from etl.irpf import extract as irpf

# Colunas que o extrator lê de cada aba (na ordem em que aparecem na planilha da Receita).
COLS8 = ["tipo_formulario", "uf", irpf.FAIXA_COL, "qtde_contribuintes", "rendimento_tributavel_total",
         "rend_sujeitos_a_tribut_exclusiva", "rend_isentos_e_nao_tributaveis", "rend_isentos_lucros_dividendos",
         "rend_isentos_socio_me_simples", "rend_isentos_transf_heranca", "rend_isentos_aposentadoria_65",
         "rend_isentos_atividade_rural", "rend_isentos_indenizacao_fgts", "rend_isentos_poupanca",
         "rend_isentos_incorp_res_cap_bonif_acao", "rend_isentos_outros", "imposto_devido_total", "bens_e_direitos"]
COLS_OC = ["tipo_formulario", "uf", "natureza_da_ocupacao", "qtde_contribuintes", "rendimento_tributavel_total",
           "rend_sujeitos_a_tribut_exclusiva", "rend_isentos_e_nao_tributaveis", "rend_isentos_lucros_dividendos",
           "rend_isentos_socio_me_simples", "imposto_devido_total"]


def _aba(wb, nome: str, colunas: list[str], linhas: list[list]):
    sh = wb.create_sheet(nome)
    sh.append([f"Tabela de teste — {nome}"])
    sh.append(colunas)
    for l in linhas:
        sh.append(l)


def _planilha() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    # faixa, contribuintes, tributável, exclusiva, isentos, dividendos, sócio ME, herança, …, imposto, bens
    _aba(wb, "Tab8", COLS8, [
        ["Completo", "Acre", "De 10 a 15", 30_000_000, 600e9, 100e9, 300e9, 60e9, 0, 40e9, 0, 0, 0, 0, 0, 200e9, 100e9, 5e12],
        ["Completo", "Acre", "Mais de 320", 11_000_000, 100e9, 100e9, 800e9, 400e9, 0, 300e9, 0, 0, 0, 0, 0, 100e9, 20e9, 9e12],
        ["Simplificado", "Acre", "De 10 a 15", "x", "x", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0],
    ])
    _aba(wb, "Tab15", COLS_OC, [
        ["Completo", "Acre", "Membro ou servidor público da administração direta federal", 20_000_000, 60e9, 10e9, 10e9, 1e9, 0, 8e9],
        ["Completo", "Acre", "Proprietário de empresa ou de firma individual ou empregador-titular", 21_000_000, 10e9, 10e9, 100e9, 45e9, 5e9, 4e9],
    ])
    _aba(wb, "Tab16", [*COLS_OC[:2], "ocupacao_principal", *COLS_OC[3:]], [
        ["Completo", "Acre", "Médico", 30_000_000, 60e9, 10e9, 30e9, 6e9, 0, 10e9],
        ["Completo", "Acre", None, 11_000_000, 10e9, 10e9, 80e9, 40e9, 0, 2e9],
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cliente_sgs(valor: float = 1412.0) -> httpx.Client:
    def handler(request: httpx.Request) -> httpx.Response:
        if "bcdata.sgs.1619" in str(request.url):
            return httpx.Response(200, json=[{"data": "01/01/2024", "valor": "1320.00"},
                                             {"data": "01/12/2024", "valor": f"{valor:.2f}"}])
        return httpx.Response(404, json={})
    return httpx.Client(transport=httpx.MockTransport(handler))


def _dados() -> dict:
    return irpf.montar(_planilha(), _cliente_sgs(), 2025, 2024, "http://exemplo/planilha.xlsx")


def test_salario_minimo_vem_do_banco_central_e_converte_as_faixas():
    dados = _dados()
    assert dados["salario_minimo"]["mensal"] == 1412.0
    assert dados["salario_minimo"]["anual"] == 16944.0
    f = next(f for f in dados["faixas"] if f["faixa"] == "De 10 a 15")
    assert f["piso_anual"] == pytest.approx(169_440) and f["teto_anual"] == pytest.approx(254_160)
    assert dados["faixas"][-1]["teto_anual"] is None          # faixa aberta no topo


def test_faixas_saem_na_ordem_da_receita_e_nao_na_alfabetica():
    dados = _dados()
    presentes = [f["faixa"] for f in dados["faixas"]]
    assert presentes == ["De 10 a 15", "Mais de 320"]
    assert dados["faixas"][0]["piso_sm"] < dados["faixas"][1]["piso_sm"]


def test_celula_suprimida_conta_como_zero_e_e_registrada():
    dados = _dados()
    assert dados["celulas_suprimidas"] == 2                   # as duas células "x" da terceira linha
    assert dados["declarantes"] == 41_000_000                 # a linha suprimida não inventa contribuintes


def test_aliquota_efetiva_e_parcela_isenta():
    dados = _dados()
    topo = dados["faixas"][-1]
    assert topo["renda"] == pytest.approx(1e12)               # 100 + 100 + 800 bi
    assert topo["aliquota_efetiva"] == pytest.approx(0.02)    # 20 bi de imposto sobre 1 tri
    assert topo["pct_isenta"] == pytest.approx(0.8)
    assert topo["isentos"]["Lucros e dividendos"] == pytest.approx(400e9)
    assert dados["aliquota_media"] == pytest.approx(120e9 / 2000e9)


def test_pesos_somam_um_entre_as_faixas():
    dados = _dados()
    assert sum(f["pct_da_renda"] for f in dados["faixas"]) == pytest.approx(1.0, abs=1e-3)
    assert sum(f["pct_do_imposto"] for f in dados["faixas"]) == pytest.approx(1.0, abs=1e-3)


def test_ocupacao_nula_vira_nao_informada_e_nao_desaparece():
    dados = _dados()
    nomes = {o["nome"] for o in dados["ocupacoes"]}
    assert nomes == {"Médico", "Não informada"}
    assert sum(o["contribuintes"] for o in dados["ocupacoes"]) == dados["declarantes"]


def test_natureza_recebe_grupo_para_a_tabela():
    dados = _dados()
    grupos = {n["nome"].split()[0]: n["grupo"] for n in dados["natureza_ocupacao"]}
    assert grupos["Membro"] == "Setor público"
    assert grupos["Proprietário"] == "Donos de negócio"


def test_comparacao_com_o_subsidio_usa_o_valor_do_ano_calendario():
    dados = _dados()
    sub = dados["subsidio_deputado"]
    assert sub["ano"] == 2024
    assert sub["mensal_dezembro"] == pytest.approx(44_008.52)      # valor vigente em dez/2024
    assert sub["anual"] == pytest.approx(584_422.67, abs=1)        # 12 meses + 13º + 1/3, pelo valor de cada mês
    assert sub["em_sm"] == pytest.approx(34.5, abs=0.1)
    assert sub["faixa"] == "De 30 a 40" or sub["faixa"] is None    # faixa ausente nesta pasta sintética


def test_totais_divergentes_entre_abas_derrubam_o_run():
    wb = openpyxl.load_workbook(io.BytesIO(_planilha()))
    wb["Tab16"].cell(row=3, column=4).value = 1_000_000            # Tab16 passa a somar menos que Tab8
    buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(SystemExit, match="totais divergem"):
        irpf.montar(buf.getvalue(), _cliente_sgs(), 2025, 2024, "http://exemplo/planilha.xlsx")
